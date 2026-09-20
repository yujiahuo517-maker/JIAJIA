#!/usr/bin/env python3
import argparse
import csv
import json
import re
from collections import OrderedDict
from pathlib import Path


ALIASES = {"spu", "spuid", "spu编码", "spu编号", "商品编码", "商品编号", "productid"}


def normalize_header(value):
    return re.sub(r"[\s_-]+", "", str(value or "").strip().lower())


def normalize_id(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not value.is_integer():
            return None
        return str(int(value))
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    return text if re.fullmatch(r"\d+", text) else None


def rows_from_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
        except csv.Error:
            dialect = csv.excel_tab if "\t" in sample else csv.excel
        yield from csv.reader(handle, dialect)


def rows_from_xlsx(path, sheet_name):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("读取 xlsx 需要 openpyxl：python -m pip install openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    for row in worksheet.iter_rows(values_only=True):
        yield list(row)


def read_spus(path, sheet_name=None, column_name=None):
    suffix = path.suffix.lower()
    rows = list(rows_from_xlsx(path, sheet_name) if suffix in {".xlsx", ".xlsm"} else rows_from_csv(path))
    if not rows:
        raise SystemExit("底表为空")
    headers = [normalize_header(value) for value in rows[0]]
    wanted = normalize_header(column_name) if column_name else None
    candidates = [index for index, header in enumerate(headers) if header == wanted] if wanted else [index for index, header in enumerate(headers) if header in ALIASES]
    if len(candidates) != 1:
        raise SystemExit(f"无法唯一识别 SPU 列；表头={rows[0]!r}，请用 --column 指定")
    column_index = candidates[0]
    values = OrderedDict()
    invalid = []
    for row_number, row in enumerate(rows[1:], start=2):
        raw = row[column_index] if column_index < len(row) else None
        if raw in (None, ""):
            continue
        value = normalize_id(raw)
        if value is None:
            invalid.append({"row": row_number, "value": str(raw)})
        else:
            values.setdefault(value, row_number)
    return list(values), invalid


def load_resolved(path):
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    rows = data.get("rows", data) if isinstance(data, dict) else data
    if not isinstance(rows, list):
        raise SystemExit("--resolved 必须是 JSON 数组或包含 rows 数组的对象")
    normalized = []
    for index, row in enumerate(rows, start=1):
        if not isinstance(row, dict):
            raise SystemExit(f"resolved 第 {index} 行不是对象")
        spu = normalize_id(row.get("spuId", row.get("spu_id", row.get("productId"))))
        sku = normalize_id(row.get("skuId", row.get("sku_id")))
        category = normalize_id(row.get("categoryId", row.get("category_id", row.get("categoryId3"))))
        if not spu or not sku or not category:
            raise SystemExit(f"resolved 第 {index} 行缺少有效 spuId/skuId/categoryId")
        normalized.append({"spuId": spu, "skuId": sku, "categoryId": category})
    return normalized


def make_batches(source_spus, resolved, max_skus):
    source_order = {spu: index for index, spu in enumerate(source_spus)}
    deduped = OrderedDict()
    for row in resolved:
        if row["spuId"] in source_order:
            deduped.setdefault((row["spuId"], row["skuId"]), row)
    grouped = {}
    for row in deduped.values():
        grouped.setdefault(row["categoryId"], []).append(row)
    batches = []
    for category_id, rows in grouped.items():
        rows.sort(key=lambda item: (source_order[item["spuId"]], item["skuId"]))
        for offset in range(0, len(rows), max_skus):
            items = rows[offset : offset + max_skus]
            batches.append({
                "batch": len(batches) + 1,
                "categoryId": category_id,
                "spuCount": len({item["spuId"] for item in items}),
                "skuCount": len(items),
                "skuList": items,
            })
    matched = {row["spuId"] for row in deduped.values()}
    batch_membership = {}
    for batch in batches:
        for item in batch["skuList"]:
            batch_membership.setdefault(item["spuId"], set()).add(batch["batch"])
    split_spus = [spu for spu in source_spus if len(batch_membership.get(spu, set())) > 1]
    return batches, [spu for spu in source_spus if spu not in matched], split_spus


def main():
    parser = argparse.ArgumentParser(description="读取品质标签 SPU 底表并按三级类目/100 SKU 生成批次")
    parser.add_argument("source", type=Path)
    parser.add_argument("--sheet")
    parser.add_argument("--column")
    parser.add_argument("--resolved", type=Path, help="含 spuId/skuId/categoryId 的接口解析结果 JSON")
    parser.add_argument("--max-skus", type=int, default=100)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if not 1 <= args.max_skus <= 100:
        raise SystemExit("--max-skus 必须在 1 到 100 之间")
    source_spus, invalid = read_spus(args.source, args.sheet, args.column)
    result = {
        "source": str(args.source.resolve()),
        "sourceSpus": source_spus,
        "sourceSpuCount": len(source_spus),
        "invalidRows": invalid,
    }
    if args.resolved:
        batches, unresolved, split_spus = make_batches(source_spus, load_resolved(args.resolved), args.max_skus)
        result.update({
            "batches": batches,
            "batchCount": len(batches),
            "resolvedSkuCount": sum(item["skuCount"] for item in batches),
            "unresolvedSpus": unresolved,
            "splitSpus": split_spus,
        })
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text + "\n", encoding="utf-8")
    print(text)


if __name__ == "__main__":
    main()
